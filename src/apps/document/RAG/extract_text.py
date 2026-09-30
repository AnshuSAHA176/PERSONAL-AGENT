import csv
import pymupdf

from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

from paddleocr import PaddleOCR
from faster_whisper import WhisperModel


# =========================================================
# INITIALIZE MODELS ONCE
# =========================================================

ocr = None
whisper = None


def get_ocr():

    global ocr

    if ocr is None:
        ocr = PaddleOCR(lang="en")

    return ocr


whisper = None


def get_whisper():

    global whisper

    if whisper is None:
        whisper = WhisperModel(
            "tiny",
            device="cpu",
            compute_type="int8"
        )

    return whisper


# =========================================================
# PDF
# =========================================================

def extract_pdf(document_path):

    doc = pymupdf.open(document_path)

    text = []

    for page_number, page in enumerate(doc, start=1):

        page_text = page.get_text("text").strip()

        if page_text:

            text.append({
                "text": page_text,
                "metadata": {
                    "page": page_number
                }
            })

    doc.close()

    return text


# =========================================================
# TXT
# =========================================================

def extract_txt(document_path):

    with open(
        document_path,
        "r",
        encoding="utf-8"
    ) as file:

        content = file.read().strip()

    if not content:
        return []

    return [
        {
            "text": content,
            "metadata": {}
        }
    ]


# =========================================================
# DOCX
# =========================================================

def extract_docx(document_path):

    doc = Document(document_path)

    text = []

    for paragraph_number, paragraph in enumerate(
        doc.paragraphs,
        start=1
    ):

        paragraph_text = paragraph.text.strip()

        if paragraph_text:

            text.append({
                "text": paragraph_text,
                "metadata": {
                    "paragraph": paragraph_number
                }
            })

    return text


# =========================================================
# PPTX
# =========================================================

def extract_pptx(document_path):

    presentation = Presentation(document_path)

    text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                shape_text = shape.text.strip()

                if shape_text:
                    slide_text.append(shape_text)

        combined_text = "\n".join(slide_text).strip()

        if combined_text:

            text.append({
                "text": combined_text,
                "metadata": {
                    "slide": slide_number
                }
            })

    return text


# =========================================================
# XLSX
# =========================================================

def extract_xlsx(document_path):

    workbook = load_workbook(
        document_path,
        read_only=True,
        data_only=True
    )

    text = []

    for worksheet in workbook.worksheets:

        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True),
            start=1
        ):

            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if not values:
                continue

            row_text = " | ".join(values)

            text.append({
                "text": row_text,
                "metadata": {
                    "sheet": worksheet.title,
                    "row": row_number
                }
            })

    workbook.close()

    return text


# =========================================================
# CSV
# =========================================================

def extract_csv(document_path):

    text = []

    with open(
        document_path,
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        reader = csv.reader(file)

        for row_number, row in enumerate(
            reader,
            start=1
        ):

            values = [
                value.strip()
                for value in row
                if value.strip()
            ]

            if not values:
                continue

            row_text = " | ".join(values)

            text.append({
                "text": row_text,
                "metadata": {
                    "row": row_number
                }
            })

    return text


# =========================================================
# IMAGE / OCR
# =========================================================

def extract_image(document_path):

    ocr_model = get_ocr()

    result = ocr_model.predict(document_path)


    text = []

    for page in result:

        data = page.json

        if not isinstance(data, dict):
            continue

        # PaddleOCR output can differ between versions.
        # Extract the recognized text from the returned structure.

        if "rec_texts" in data:

            for recognized_text in data["rec_texts"]:

                recognized_text = recognized_text.strip()

                if recognized_text:

                    text.append({
                        "text": recognized_text,
                        "metadata": {}
                    })

    return text


# =========================================================
# AUDIO / WHISPER
# =========================================================

def extract_audio(document_path):

    whisper_model = get_whisper()

    segments, info = whisper_model.transcribe(
        document_path,
        beam_size=5,
        vad_filter=True
    )


    text = []

    for segment in segments:

        segment_text = segment.text.strip()

        if not segment_text:
            continue

        text.append({
            "text": segment_text,
            "metadata": {
                "start": segment.start,
                "end": segment.end,
                "language": info.language
            }
        })

    return text


# =========================================================
# MAIN DOCUMENT EXTRACTOR
# =========================================================

def extract_document(file_format, document_path):

    file_format = file_format.lower()

    if file_format == ".pdf":

        return extract_pdf(document_path)

    elif file_format == ".txt":

        return extract_txt(document_path)

    elif file_format == ".docx":

        return extract_docx(document_path)

    elif file_format == ".pptx":

        return extract_pptx(document_path)

    elif file_format == ".xlsx":

        return extract_xlsx(document_path)

    elif file_format == ".csv":

        return extract_csv(document_path)

    # elif file_format in [
    #     ".jpg",
    #     ".jpeg",
    #     ".png",
    #     ".webp"
    # ]:

    #     return extract_image(document_path)

    elif file_format in [
        ".mp3",
        ".wav",
        ".m4a",
        ".flac",
        ".ogg"
    ]:

        return extract_audio(document_path)

    else:

        raise ValueError(
            f"Unsupported file format: {file_format}"
        )